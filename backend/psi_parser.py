import re
from pathlib import Path

WEEK_RE = re.compile(r'^W\d{2}$')


def parse_psi_sheet(path: str) -> tuple[dict, list[str]]:
    ext = Path(path).suffix.lower()
    if ext == '.xlsb':
        rows = _read_xlsb(path)
    else:
        rows = _read_xlsx(path)
    return _process_rows(rows)


def _read_xlsx(path: str) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = _find_sheet(wb.sheetnames)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def _read_xlsb(path: str) -> list:
    from pyxlsb import open_workbook
    rows = []
    with open_workbook(path) as wb:
        sheet_name = _find_sheet(wb.sheets)
        with wb.get_sheet(sheet_name) as ws:
            for row in ws.rows():
                if not row:
                    rows.append(())
                    continue
                # Cell.c is 0-based column index
                max_col = max(c.c for c in row)
                row_vals = [None] * (max_col + 1)
                for cell in row:
                    row_vals[cell.c] = cell.v
                rows.append(tuple(row_vals))
    return rows


def _find_sheet(sheet_names) -> str:
    for name in sheet_names:
        if 'PSI' in name.upper():
            return name
    return sheet_names[0]


def _process_rows(rows: list) -> tuple[dict, list[str]]:
    if len(rows) < 3:
        return {}, []

    header = rows[2]  # row index 2 = row 3 (main header)

    # Collect all columns from index 11 (col L) onward that match W\d{2}
    week_headers = []
    week_col_indices = []
    for i in range(11, len(header)):
        h = header[i]
        if h and WEEK_RE.match(str(h).strip()):
            week_headers.append(str(h).strip())
            week_col_indices.append(i)

    data_map = {}
    for row in rows[3:]:
        if not row or len(row) < 11 or not row[1]:
            continue
        item = str(row[1]).strip()
        if not item:
            continue
        seg_type = str(row[10]).strip() if row[10] is not None else ""
        if not seg_type:
            continue

        weeks = [_num(row[wi]) if wi < len(row) else 0.0 for wi in week_col_indices]

        key = f"{item}|{seg_type}"
        data_map[key] = {
            "item": item,
            "vendor": str(row[3]).strip() if row[3] is not None else "",
            "ss":     _num(row[5]) if len(row) > 5 else 0.0,
            "ending": _num(row[6]) if len(row) > 6 else 0.0,
            "desc":   str(row[7]).strip() if len(row) > 7 and row[7] is not None else "",
            "moq":    _num(row[8]) if len(row) > 8 else 0.0,
            "lt":     _num(row[9]) if len(row) > 9 else 0.0,
            "seg_type": seg_type,
            "weeks": weeks,
        }

    # Group by item code to collect full metadata and share it
    item_metadata = {}
    for key, data in data_map.items():
        item = data["item"]
        if item not in item_metadata:
            item_metadata[item] = {
                "vendor": "",
                "desc": "",
                "moq": 0.0,
                "lt": 0.0,
                "ss": 0.0,
                "ending": 0.0
            }
        
        # Capture non-empty/non-zero metadata values
        if data["vendor"]:
            item_metadata[item]["vendor"] = data["vendor"]
        if data["desc"]:
            item_metadata[item]["desc"] = data["desc"]
        if data["moq"] > 0:
            item_metadata[item]["moq"] = data["moq"]
        if data["lt"] > 0:
            item_metadata[item]["lt"] = data["lt"]
        if data["ss"] > 0:
            item_metadata[item]["ss"] = data["ss"]
        if data["ending"] > 0:
            item_metadata[item]["ending"] = data["ending"]

    # Share collected metadata back to all rows of the same item
    for key, data in data_map.items():
        item = data["item"]
        meta = item_metadata.get(item)
        if meta:
            if meta["vendor"]:
                data["vendor"] = meta["vendor"]
            if meta["desc"]:
                data["desc"] = meta["desc"]
            if meta["moq"] > 0:
                data["moq"] = meta["moq"]
            if meta["lt"] > 0:
                data["lt"] = meta["lt"]
            if meta["ss"] > 0:
                data["ss"] = meta["ss"]
            if meta["ending"] > 0:
                data["ending"] = meta["ending"]

    return data_map, week_headers


def _num(val) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0
